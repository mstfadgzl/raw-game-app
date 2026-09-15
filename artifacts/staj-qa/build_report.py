from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION_START
from copy import deepcopy
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import textwrap
from report_content import days

ROOT=Path('C:/MyProject/RawgGameApp')
QA=ROOT/'artifacts/staj-qa'
OUT=ROOT/'artifacts/GameCodex_20_Gunluk_Staj_Raporu.docx'
figures=QA/'figures'; figures.mkdir(exist_ok=True)
font=ImageFont.truetype('C:/Windows/Fonts/consola.ttf',25)
small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24)

def code_figure(day,path,anchor,n):
    lines=(ROOT/path).read_text(encoding='utf-8-sig').splitlines()
    start=next(i for i,line in enumerate(lines) if anchor in line)
    selected=lines[start:start+n]
    indent=min(len(s)-len(s.lstrip()) for s in selected if s.strip())
    display=[]
    for k,line in enumerate(selected):
        pieces=textwrap.wrap(line[indent:],width=74,replace_whitespace=False,drop_whitespace=False) or ['']
        display.extend([(str(start+k+1) if j==0 else '',t) for j,t in enumerate(pieces)])
    im=Image.new('RGB',(1280,74+len(display)*33+20),'white')
    d=ImageDraw.Draw(im)
    d.text((16,10),path,font=small,fill='black')
    d.line((16,49,1264,49),fill='#bbbbbb',width=2)
    for j,(num,line) in enumerate(display):
        d.text((16,65+j*33),num,font=font,fill='#777777')
        d.text((85,65+j*33),line,font=font,fill='#151515')
    target=figures/f'figure-{day:02}.png';im.save(target)
    return target,start+1,start+len(selected)

doc=Document(QA/'reference.docx')
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'): body.remove(child)
section=doc.sections[0]
section.top_margin=Cm(5.0)
section.bottom_margin=Cm(5.0)
normal=doc.styles['Normal']
normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.space_before=Pt(0)
normal.paragraph_format.widow_control=True
for name in ['Title','Heading 1','Caption']:
    if name not in doc.styles:
        doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH).base_style=normal
    s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(12 if name!='Caption' else 9)
    s.paragraph_format.space_after=Pt(8 if name!='Caption' else 0)
    s.paragraph_format.space_before=Pt(0)

# Preserve supplied institutional form, filling only its blank date and page slots.
table=section.header.tables[0]
p=table.cell(0,3).paragraphs[0]
p.add_run('… / … / 20…').font.size=Pt(10)
p=table.cell(1,3).paragraphs[0]
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);field.append(r);p._p.append(field)

for i,(title,paras,path,anchor,n) in enumerate(days,1):
    if i>1:
        previous=doc.sections[-1].header
        old_part=previous.part
        old_nodes=[deepcopy(c) for c in previous._element]
        new_section=doc.add_section(WD_SECTION_START.NEW_PAGE)
        new_section.header.is_linked_to_previous=False
        header=new_section.header
        for c in list(header._element): header._element.remove(c)
        for c in old_nodes: header._element.append(c)
        for rid,rel in old_part.rels.items():
            newrid=header.part.relate_to(rel.target_part,rel.reltype)
            for elem in header._element.iter():
                if elem.get(qn('r:embed'))==rid: elem.set(qn('r:embed'),newrid)
    p=doc.add_paragraph(style='Title' if i==1 else 'Heading 1')
    p.paragraph_format.page_break_before=False
    p.paragraph_format.keep_with_next=True
    r=p.add_run(f'{i}. Gün  {title}');r.bold=True
    for para in paras:
        p=doc.add_paragraph(para)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(4)
    if path=='SCREENSHOT':
        picture=Path('C:/Users/musta/AppData/Local/Temp/gamecodex-stores.png')
        p.add_run().add_picture(str(picture),height=Inches(3.55))
        caption=f'Şekil {i}. Gerçek Android cihazında mağaza seçim ekranı.'
    elif path=='ICON':
        p.add_run().add_picture(str(ROOT/'app/assets/brand/gamecodex.png'),height=Cm(4.0))
        caption=f'Şekil {i}. GameCodex uygulaması için kullanılan simge.'
    else:
        picture,start,end=code_figure(i,path,anchor,n)
        p.add_run().add_picture(str(picture),width=Cm(16.0))
        caption=f'Şekil {i}. Projedeki koddan kesit, {start}–{end}. satırlar.'
    cp=doc.add_paragraph(caption,style='Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next=False
    print(i,len(' '.join(paras).split()),title)

doc.core_properties.title='GameCodex 20 Günlük Staj Raporu'
doc.core_properties.subject='Mobil oyun keşif uygulamasının geliştirilmesi'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.comments=''
doc.save(OUT)
print(OUT)
